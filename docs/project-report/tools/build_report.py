"""Builds AI_Career_Companion_Project_Report.docx from the Markdown source.

Run from the repository root with the project virtual environment (python-docx is a
backend dependency):

    .venv\\Scripts\\python.exe docs\\project-report\\tools\\build_report.py

The Markdown subset understood here is the one used by the report: chapter / section
headings with optional {.prelim} {.unnumbered} {.appendix} {.title} {.center}
attributes, paragraphs with **bold**, *italic*, `code` and $math$, bullet and numbered
lists, pipe tables preceded by a "Table X.Y: caption" line, ![Figure X.Y: caption](path)
images, fenced code blocks, $$ display math $$, > notes, and HTML-comment directives
(titlepage, frontmatter, mainmatter, toc, lof, lot, signatures, references).

Table of contents, list of figures and list of tables are Word fields; they are filled
in by export_pdf.ps1 (Word updates the fields, saves the DOCX and exports the PDF).
"""

from __future__ import annotations

import re
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "AI_Career_Companion_Project_Report.md"
OUTPUT = ROOT / "AI_Career_Companion_Project_Report.docx"

BODY_FONT = "Times New Roman"
CODE_FONT = "Consolas"
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
GREY = RGBColor(0x55, 0x5F, 0x6B)
HEADER_TEXT = "AI Career Companion — Academic Project Report"
TEXT_WIDTH_IN = 8.27 - 1.5 - 1.0  # A4 width minus left and right margins


# ---------------------------------------------------------------------------
# Low-level XML helpers
# ---------------------------------------------------------------------------
def set_cell_shading(cell, hex_fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def set_paragraph_shading(paragraph, hex_fill: str, border: str | None = None) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    if border:
        bdr = OxmlElement("w:pBdr")
        for side in ("top", "left", "bottom", "right"):
            el = OxmlElement(f"w:{side}")
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "4")
            el.set(qn("w:space"), "4")
            el.set(qn("w:color"), border)
            bdr.append(el)
        p_pr.append(bdr)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    p_pr.append(shd)


def add_bottom_border(paragraph, color="8A94A0") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    bdr.append(bottom)
    p_pr.append(bdr)


def add_field(paragraph, instruction: str, placeholder: str = "") -> None:
    """Inserts a complex field (TOC, PAGE, ...) into a paragraph."""
    def run_with(child):
        r = OxmlElement("w:r")
        r.append(child)
        paragraph._p.append(r)
        return r

    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    run_with(begin)
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruction} "
    run_with(instr)
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    run_with(sep)
    if placeholder:
        t = OxmlElement("w:t")
        t.text = placeholder
        run_with(t)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run_with(end)


def set_page_numbering(section, fmt: str | None, start: int | None) -> None:
    sect_pr = section._sectPr
    for old in sect_pr.findall(qn("w:pgNumType")):
        sect_pr.remove(old)
    pg = OxmlElement("w:pgNumType")
    if fmt:
        pg.set(qn("w:fmt"), fmt)
    if start is not None:
        pg.set(qn("w:start"), str(start))
    sect_pr.append(pg)


def set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def set_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def set_table_borders(table, color="9AA3AD") -> None:
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tbl_pr.append(borders)


def set_cell_margins(table, top=40, bottom=40, left=80, right=80) -> None:
    tbl_pr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for side, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tbl_pr.append(mar)


def set_font(style_or_run, name: str) -> None:
    style_or_run.font.name = name
    r_pr = style_or_run.element.get_or_add_rPr()
    fonts = r_pr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        r_pr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), name)


# ---------------------------------------------------------------------------
# Inline formatting and a tiny LaTeX-to-text converter for the report's formulas
# ---------------------------------------------------------------------------
LATEX_SYMBOLS = {
    r"\cdot": "·", r"\times": "×", r"\ge": "≥", r"\le": "≤", r"\in": "∈", r"\sum": "Σ",
    r"\dots": "…", r"\cos": "cos", r"\{": "{", r"\}": "}", r"\!": "", r"\,": " ", r"\left": "",
    r"\right": "",
}


def latex_runs(src: str) -> list[tuple[str, dict]]:
    """Converts the small LaTeX subset used in the report to (text, format) runs."""
    out: list[tuple[str, dict]] = []
    i = 0

    def read_group(j: int) -> tuple[str, int]:
        if j < len(src) and src[j] == "{":
            depth, k = 1, j + 1
            while k < len(src) and depth:
                depth += {"{": 1, "}": -1}.get(src[k], 0)
                k += 1
            return src[j + 1:k - 1], k
        if src.startswith("\\", j):
            m = re.match(r"\\[a-zA-Z]+", src[j:])
            if m:
                return m.group(0), j + len(m.group(0))
        return src[j:j + 1], j + 1

    while i < len(src):
        ch = src[i]
        if src.startswith(r"\mathbf", i):
            body, i = read_group(i + len(r"\mathbf"))
            out.extend((t, {**f, "bold": True}) for t, f in latex_runs(body))
            continue
        if src.startswith(r"\mathrm", i):
            body, i = read_group(i + len(r"\mathrm"))
            out.extend((t, {**f, "upright": True}) for t, f in latex_runs(body))
            continue
        if src.startswith(r"\frac", i):
            num, i = read_group(i + len(r"\frac"))
            den, i = read_group(i)
            out.append(("(", {}))
            out.extend(latex_runs(num))
            out.append((") / (", {}))
            out.extend(latex_runs(den))
            out.append((")", {}))
            continue
        if ch in "_^":
            body, i = read_group(i + 1)
            key = "subscript" if ch == "_" else "superscript"
            out.extend((t, {**f, key: True}) for t, f in latex_runs(body))
            continue
        if ch == "\\":
            m = re.match(r"\\[a-zA-Z]+|\\[{}!,]", src[i:])
            token = m.group(0) if m else ch
            out.append((LATEX_SYMBOLS.get(token, token.lstrip("\\")), {"upright": token in (r"\cos",)}))
            i += len(token)
            continue
        if ch == " ":
            out.append((" ", {}))
            i += 1
            continue
        out.append((ch, {}))
        i += 1
    merged: list[tuple[str, dict]] = []
    for text, fmt in out:
        if merged and merged[-1][1] == fmt:
            merged[-1] = (merged[-1][0] + text, fmt)
        else:
            merged.append((text, fmt))
    return merged


INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\$[^$]+\$|\*[^*\s][^*]*?\*)")


def add_inline(paragraph, text: str, base: dict | None = None) -> None:
    base = base or {}
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            add_inline(paragraph, part[2:-2], {**base, "bold": True})
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_font(run, CODE_FONT)
            run.font.size = Pt(10)
            run.bold = base.get("bold")
        elif part.startswith("$") and part.endswith("$") and len(part) > 2:
            for t, f in latex_runs(part[1:-1]):
                run = paragraph.add_run(t)
                run.italic = not f.get("upright") and not f.get("bold") and any(c.isalpha() for c in t) and len(t) < 4
                run.bold = f.get("bold") or base.get("bold")
                run.font.subscript = f.get("subscript")
                run.font.superscript = f.get("superscript")
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            add_inline(paragraph, part[1:-1], {**base, "italic": True})
        else:
            run = paragraph.add_run(part)
            run.bold = base.get("bold")
            run.italic = base.get("italic")


# ---------------------------------------------------------------------------
# Document setup
# ---------------------------------------------------------------------------
def setup_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    set_font(normal, BODY_FONT)
    normal.font.size = Pt(12)
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_after = Pt(6)
    pf.space_before = Pt(0)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for name, size, before, after in (("Heading 1", 16, 0, 18), ("Heading 2", 13.5, 14, 6), ("Heading 3", 12, 10, 4)):
        st = styles[name]
        set_font(st, BODY_FONT)
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.italic = name == "Heading 3"
        st.font.color.rgb = NAVY
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True
        st.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    styles["Heading 1"].paragraph_format.page_break_before = True

    def para_style(name, base="Normal"):
        st = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        st.base_style = styles[base]
        return st

    prelim = para_style("Prelim Heading")
    prelim.font.size = Pt(15)
    prelim.font.bold = True
    prelim.font.color.rgb = NAVY
    prelim.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    prelim.paragraph_format.space_after = Pt(18)
    prelim.paragraph_format.page_break_before = True
    prelim.paragraph_format.keep_with_next = True
    prelim.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    prelim_pr = prelim.element.get_or_add_pPr()
    outline = OxmlElement("w:outlineLvl")
    outline.set(qn("w:val"), "0")
    prelim_pr.append(outline)

    notoc = para_style("Prelim Heading NoTOC", "Prelim Heading")
    notoc_pr = notoc.element.get_or_add_pPr()
    for el in notoc_pr.findall(qn("w:outlineLvl")):
        notoc_pr.remove(el)

    for name in ("Figure Caption", "Table Caption"):
        cap = para_style(name)
        cap.font.size = Pt(11)
        cap.font.bold = name == "Table Caption"
        cap.font.italic = name == "Figure Caption"
        cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        cap.paragraph_format.space_before = Pt(4 if name == "Figure Caption" else 10)
        cap.paragraph_format.space_after = Pt(12 if name == "Figure Caption" else 4)
        cap.paragraph_format.keep_with_next = name == "Table Caption"

    code = para_style("Code Block")
    set_font(code, CODE_FONT)
    code.font.size = Pt(9)
    code.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    code.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    code.paragraph_format.space_before = Pt(4)
    code.paragraph_format.space_after = Pt(10)
    code.paragraph_format.left_indent = Cm(0.2)
    code.paragraph_format.right_indent = Cm(0.2)

    cell = para_style("Table Text")
    cell.font.size = Pt(10)
    cell.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    cell.paragraph_format.space_after = Pt(0)
    cell.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    note = para_style("Note")
    note.font.size = Pt(11)
    note.font.italic = True
    note.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    note.paragraph_format.space_after = Pt(12)
    note.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    eq = para_style("Equation")
    eq.font.size = Pt(12)
    eq.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    eq.paragraph_format.space_before = Pt(4)
    eq.paragraph_format.space_after = Pt(8)

    ref = para_style("Reference")
    ref.font.size = Pt(11)
    ref.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    ref.paragraph_format.left_indent = Cm(1.0)
    ref.paragraph_format.first_line_indent = Cm(-1.0)
    ref.paragraph_format.space_after = Pt(6)
    ref.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    for name in ("List Bullet",):
        st = styles[name]
        st.paragraph_format.space_after = Pt(3)
        st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    numbered = para_style("Numbered Item")
    numbered.paragraph_format.left_indent = Cm(0.9)
    numbered.paragraph_format.first_line_indent = Cm(-0.6)
    numbered.paragraph_format.space_after = Pt(3)
    numbered.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    for toc_level in ("TOC 1", "TOC 2", "TOC 3"):
        try:
            st = styles[toc_level]
        except KeyError:
            st = styles.add_style(toc_level, WD_STYLE_TYPE.PARAGRAPH)
        st.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        st.paragraph_format.space_after = Pt(3)


def configure_section(section) -> None:
    section.page_width = Emu(int(8.27 * 914400))
    section.page_height = Emu(int(11.69 * 914400))
    section.left_margin = Inches(1.5)
    section.right_margin = Inches(1.0)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.header_distance = Inches(0.5)
    section.footer_distance = Inches(0.5)


def page_footer(section) -> None:
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(p, "PAGE", "1")
    for run in p.runs:
        run.font.size = Pt(10)


def page_header(section, text: str | None) -> None:
    header = section.header
    header.is_linked_to_previous = False
    p = header.paragraphs[0]
    p.text = ""
    if text:
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = p.add_run(text)
        run.font.size = Pt(9)
        run.italic = True
        run.font.color.rgb = GREY
        add_bottom_border(p)


# ---------------------------------------------------------------------------
# Block builders
# ---------------------------------------------------------------------------
def add_table(doc, rows: list[list[str]]) -> None:
    header, body = rows[0], rows[1:]
    ncols = len(header)
    table = doc.add_table(rows=1 + len(body), cols=ncols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    set_cell_margins(table)
    # Column widths proportional to (damped) content length, with a floor.
    lengths = []
    for c in range(ncols):
        longest = max(len(re.sub(r"[`*$]", "", r[c])) if c < len(r) else 0 for r in rows)
        lengths.append(max(6, min(longest, 70)) ** 0.75)
    total = sum(lengths)
    widths = [TEXT_WIDTH_IN * w / total for w in lengths]
    # Never narrower than the column's longest word, so short labels ("NFR-01") do not split.
    floors = []
    for c in range(ncols):
        words = [w for r in rows if c < len(r) for w in re.sub(r"[`*$]", "", r[c]).split()]
        floors.append(min(1.6, 0.16 + 0.078 * max((len(w) for w in words), default=4)))
    for _ in range(3):
        short = [c for c in range(ncols) if widths[c] < floors[c]]
        if not short:
            break
        need = sum(floors[c] - widths[c] for c in short)
        donors = [c for c in range(ncols) if c not in short]
        pool = sum(widths[c] - floors[c] for c in donors if widths[c] > floors[c])
        for c in short:
            widths[c] = floors[c]
        for c in donors:
            if widths[c] > floors[c] and pool > 0:
                widths[c] -= need * (widths[c] - floors[c]) / pool
    numeric_cols = [bool(body) and all(re.fullmatch(r"[\d.,/ %×→–\-msx@()]+", r[c].strip()) for r in body if c < len(r))
                    for c in range(ncols)]
    for r_idx, row_cells in enumerate([header] + body):
        row = table.rows[r_idx]
        set_cant_split(row)
        for c in range(ncols):
            cell = row.cells[c]
            cell.width = Inches(widths[c])
            text = row_cells[c] if c < len(row_cells) else ""
            p = cell.paragraphs[0]
            p.style = doc.styles["Table Text"]
            numeric = r_idx > 0 and numeric_cols[c]
            if numeric:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            add_inline(p, text, {"bold": True} if r_idx == 0 else {})
            if r_idx == 0:
                set_cell_shading(cell, "DCE6F1")
    set_repeat_header(table.rows[0])
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(4)
    spacer.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE


def rotated_copy(path: Path) -> Path:
    """A copy of a wide figure turned 90° counter-clockwise, for a full-page landscape placement."""
    from PIL import Image

    out = Path(tempfile.gettempdir()) / f"rotated-{path.name}"
    with Image.open(path) as im:
        im.rotate(90, expand=True).save(out)
    return out


def add_figure(doc, caption: str, path: Path, rotate: bool = False) -> None:
    from PIL import Image

    if rotate:
        path = rotated_copy(path)
    with Image.open(path) as im:
        w, h = im.size
    limit = 8.6 if rotate else 7.6  # keep each figure and its caption on one page
    width = min(TEXT_WIDTH_IN, 6.0)
    if width * h / w > limit:
        width = limit * w / h
    p = doc.add_paragraph()
    p.paragraph_format.page_break_before = rotate
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.add_run().add_picture(str(path), width=Inches(width))
    cap = doc.add_paragraph(style="Figure Caption")
    add_inline(cap, caption)


def add_signatures(doc, spec: str) -> None:
    cells = [c.strip() for c in spec.split("|")]
    pairs = [(cells[i], cells[i + 1] if i + 1 < len(cells) else "") for i in range(0, len(cells), 2)]
    table = doc.add_table(rows=1, cols=len(pairs))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (name, role) in enumerate(pairs):
        cell = table.rows[0].cells[i]
        cell.width = Inches(TEXT_WIDTH_IN / len(pairs))
        p = cell.paragraphs[0]
        p.style = doc.styles["Table Text"]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(42)
        p.add_run("_" * 26)
        p2 = cell.add_paragraph(style="Table Text")
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_inline(p2, name, {"bold": True})
        if role:
            p3 = cell.add_paragraph(style="Table Text")
            p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(p3, role)
    doc.add_paragraph()


def add_code(doc, lines: list[str]) -> None:
    p = doc.add_paragraph(style="Code Block")
    set_paragraph_shading(p, "F3F5F8", border="C8CED6")
    for i, line in enumerate(lines):
        if i:
            p.add_run().add_break(WD_BREAK.LINE)
        run = p.add_run(line.replace("\t", "    ") or " ")
        set_font(run, CODE_FONT)
        run.font.size = Pt(9)


# ---------------------------------------------------------------------------
# Main conversion
# ---------------------------------------------------------------------------
ATTR = re.compile(r"\s*\{(\.[a-z]+(?:\s+\.[a-z]+)*)\}\s*$")


def split_attrs(text: str) -> tuple[str, set[str]]:
    m = ATTR.search(text)
    if not m:
        return text, set()
    return text[:m.start()].rstrip(), {a.lstrip(".") for a in m.group(1).split()}


def parse_table(lines: list[str]) -> list[list[str]]:
    rows = []
    for line in lines:
        if re.fullmatch(r"\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)*\|?", line.strip()):
            continue
        cells = [c.strip() for c in re.split(r"(?<!\\)\|", line.strip().strip("|"))]
        rows.append([c.replace(r"\|", "|") for c in cells])
    return rows


def build() -> None:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    setup_styles(doc)
    configure_section(doc.sections[0])
    first = doc.sections[0]
    page_header(first, None)
    first.footer.is_linked_to_previous = False

    chapter = section_no = sub_no = 0
    in_title = in_refs = False
    pending_table_caption: str | None = None
    i = 0
    para_buf: list[str] = []

    def flush_para():
        nonlocal para_buf
        if not para_buf:
            return
        text = " ".join(s.strip() for s in para_buf)
        para_buf = []
        text, attrs = split_attrs(text)
        if in_refs and re.match(r"\[\d+\]", text):
            p = doc.add_paragraph(style="Reference")
            add_inline(p, text)
            return
        p = doc.add_paragraph()
        if "center" in attrs or in_title:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(4)
        add_inline(p, text)

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # Directives
        m = re.fullmatch(r"<!--\s*(.+?)\s*-->", stripped)
        if m:
            flush_para()
            directive = m.group(1)
            if directive == "titlepage":
                in_title = True
                for _ in range(3):
                    doc.add_paragraph()
            elif directive in ("frontmatter", "mainmatter"):
                in_title = False
                sec = doc.add_section(WD_SECTION.NEW_PAGE)
                configure_section(sec)
                if directive == "frontmatter":
                    set_page_numbering(sec, "lowerRoman", 1)
                    page_header(sec, None)
                else:
                    set_page_numbering(sec, "decimal", 1)
                    page_header(sec, HEADER_TEXT)
                page_footer(sec)
            elif directive == "toc":
                p = doc.add_paragraph()
                add_field(p, 'TOC \\o "1-3" \\h \\z \\t "Prelim Heading,1"', "Update fields to build the table of contents.")
            elif directive == "lof":
                p = doc.add_paragraph()
                add_field(p, 'TOC \\h \\z \\t "Figure Caption,1"', "Update fields to build the list of figures.")
            elif directive == "lot":
                p = doc.add_paragraph()
                add_field(p, 'TOC \\h \\z \\t "Table Caption,1"', "Update fields to build the list of tables.")
            elif directive.startswith("signatures:"):
                add_signatures(doc, directive.split(":", 1)[1])
            elif directive == "references":
                in_refs = True
            # "pagebreak" is implied by the page-break-before of every heading.
            i += 1
            continue

        if not stripped:
            flush_para()
            i += 1
            continue

        # Headings
        hm = re.match(r"(#{1,3})\s+(.*)", line)
        if hm:
            flush_para()
            level = len(hm.group(1))
            text, attrs = split_attrs(hm.group(2))
            if "title" in attrs:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(30)
                p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
                run = p.add_run(text)
                run.bold = True
                run.font.size = Pt(20)
                run.font.color.rgb = NAVY
            elif "prelim" in attrs:
                notoc = text.upper() in ("TABLE OF CONTENTS",)
                p = doc.add_paragraph(text, style="Prelim Heading NoTOC" if notoc else "Prelim Heading")
            elif level == 1:
                in_refs = False
                if "unnumbered" in attrs or "appendix" in attrs:
                    doc.add_paragraph(text, style="Heading 1")
                    chapter_label = None
                else:
                    chapter += 1
                    section_no = sub_no = 0
                    doc.add_paragraph(f"Chapter {chapter}: {text}", style="Heading 1")
            elif level == 2:
                section_no += 1
                sub_no = 0
                label = f"{chapter}.{section_no} " if chapter and not in_refs else ""
                doc.add_paragraph(f"{label}{text}", style="Heading 2")
            else:
                sub_no += 1
                doc.add_paragraph(f"{chapter}.{section_no}.{sub_no} {text}", style="Heading 3")
            i += 1
            continue

        # Figure
        fm = re.fullmatch(r"!\[(.+?)\]\((.+?)\)(\{\.rotate\})?", stripped)
        if fm:
            flush_para()
            add_figure(doc, fm.group(1), ROOT / fm.group(2), rotate=bool(fm.group(3)))
            i += 1
            continue

        # Table caption
        if re.match(r"Table [A-Z0-9]+\.\d+:", stripped):
            flush_para()
            pending_table_caption = stripped
            i += 1
            continue

        # Table
        if stripped.startswith("|"):
            flush_para()
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            if pending_table_caption:
                cap = doc.add_paragraph(style="Table Caption")
                add_inline(cap, pending_table_caption)
                pending_table_caption = None
            add_table(doc, parse_table(block))
            continue

        # Fenced code
        if stripped.startswith("```"):
            flush_para()
            i += 1
            block = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            i += 1
            add_code(doc, block)
            continue

        # Display math
        if stripped.startswith("$$") and stripped.endswith("$$") and len(stripped) > 4:
            flush_para()
            p = doc.add_paragraph(style="Equation")
            add_inline(p, "$" + stripped[2:-2].strip() + "$")
            i += 1
            continue

        # Note
        if stripped.startswith(">"):
            flush_para()
            p = doc.add_paragraph(style="Note")
            set_paragraph_shading(p, "FFF8E5", border="D6B25E")
            add_inline(p, stripped.lstrip("> ").strip())
            i += 1
            continue

        # Lists
        bm = re.match(r"-\s+(.*)", stripped)
        nm = re.match(r"(\d+)\.\s+(.*)", stripped)
        if bm or nm:
            flush_para()
            item = (bm or nm).group(1 if bm else 2)
            # continuation lines
            i += 1
            while i < len(lines) and lines[i].startswith("  ") and lines[i].strip() and not re.match(r"\s*(-|\d+\.)\s", lines[i]):
                item += " " + lines[i].strip()
                i += 1
            if bm:
                p = doc.add_paragraph(style="List Bullet")
            else:
                p = doc.add_paragraph(style="Numbered Item")
                p.add_run(f"{nm.group(1)}.\t")
                p.paragraph_format.tab_stops.add_tab_stop(Cm(0.9))
            add_inline(p, item)
            continue

        para_buf.append(line)
        i += 1

    flush_para()
    props = doc.core_properties
    props.title = "AI Career Companion: An Intelligent Career Recommendation and Skill-Gap Analysis Platform"
    props.subject = "Academic project report — Infosys internship project"
    props.author = "Jeswin Jarald"
    props.keywords = "career recommendation; skill-gap analysis; semantic search; FAISS; FastAPI; React"
    doc.save(OUTPUT)
    print(f"wrote {OUTPUT}")


if __name__ == "__main__":
    sys.exit(build())
