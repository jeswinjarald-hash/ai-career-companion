"""Milestone 3.2 — resume/cover-letter export.

Generates PDF (reportlab) and DOCX (python-docx) documents in memory from an already
persisted, validated `ApplicationCustomization` — never from unvalidated content, and
never written to disk (nothing here touches the resume storage directory or the
original uploaded resume file). Layout uses standard headings, a single column, and
plain paragraph/bullet flowables only, so the output stays ATS-parseable (selectable
text, no tables, no decorative graphics).
"""

import io

from docx import Document
from docx.shared import Pt
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer

from app.schemas.customization import ApplicationCustomization


def _resume_sections(customization: ApplicationCustomization) -> list[tuple[str, list[str]]]:
    resume = customization.tailored_resume
    sections: list[tuple[str, list[str]]] = []
    if resume.summary:
        sections.append(("SUMMARY", [resume.summary]))
    if resume.skills:
        sections.append(("SKILLS", [", ".join(resume.skills)]))
    if resume.education:
        sections.append(("EDUCATION", [entry.raw_text for entry in resume.education]))
    if resume.projects:
        sections.append(("PROJECTS", [f"{p.title} — {p.tailored_text}" if p.title else p.tailored_text for p in resume.projects]))
    if resume.experience:
        sections.append(("EXPERIENCE", [b.tailored_text for b in resume.experience]))
    if resume.internships:
        sections.append(("INTERNSHIPS", [b.tailored_text for b in resume.internships]))
    if resume.certifications:
        sections.append(("CERTIFICATIONS", resume.certifications))
    if resume.achievements:
        sections.append(("ACHIEVEMENTS", resume.achievements))
    return sections


def export_resume_pdf(customization: ApplicationCustomization, candidate_name: str) -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=LETTER, topMargin=0.6 * inch, bottomMargin=0.6 * inch, leftMargin=0.7 * inch, rightMargin=0.7 * inch)
    styles = getSampleStyleSheet()
    heading_style = ParagraphStyle("SectionHeading", parent=styles["Heading2"], fontSize=11, spaceBefore=10, spaceAfter=4)
    name_style = ParagraphStyle("Name", parent=styles["Title"], fontSize=16, spaceAfter=2)
    body_style = ParagraphStyle("Body", parent=styles["Normal"], fontSize=10, leading=13)

    story = [Paragraph(candidate_name or "Candidate", name_style)]
    if customization.tailored_resume.header:
        story.append(Paragraph(customization.tailored_resume.header, body_style))
    story.append(Spacer(1, 8))

    for title, items in _resume_sections(customization):
        story.append(Paragraph(title, heading_style))
        if title in ("SUMMARY", "SKILLS"):
            for item in items:
                story.append(Paragraph(item, body_style))
        else:
            story.append(ListFlowable([ListItem(Paragraph(item, body_style)) for item in items], bulletType="bullet"))

    doc.build(story)
    return buffer.getvalue()


def export_resume_docx(customization: ApplicationCustomization, candidate_name: str) -> bytes:
    document = Document()
    title = document.add_heading(candidate_name or "Candidate", level=1)
    title.runs[0].font.size = Pt(18)
    if customization.tailored_resume.header:
        document.add_paragraph(customization.tailored_resume.header)

    for section_title, items in _resume_sections(customization):
        document.add_heading(section_title, level=2)
        if section_title in ("SUMMARY", "SKILLS"):
            for item in items:
                document.add_paragraph(item)
        else:
            for item in items:
                document.add_paragraph(item, style="List Bullet")

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def export_cover_letter_pdf(customization: ApplicationCustomization, candidate_name: str) -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=LETTER, topMargin=0.75 * inch, bottomMargin=0.75 * inch, leftMargin=0.9 * inch, rightMargin=0.9 * inch)
    styles = getSampleStyleSheet()
    body_style = ParagraphStyle("Body", parent=styles["Normal"], fontSize=11, leading=16, spaceAfter=10)
    name_style = ParagraphStyle("Name", parent=styles["Title"], fontSize=14, spaceAfter=12)

    story = [Paragraph(candidate_name or "Candidate", name_style)]
    for paragraph in customization.cover_letter_text.split("\n\n"):
        if paragraph.strip():
            story.append(Paragraph(paragraph.strip(), body_style))
    doc.build(story)
    return buffer.getvalue()


def export_cover_letter_docx(customization: ApplicationCustomization, candidate_name: str) -> bytes:
    document = Document()
    document.add_heading(candidate_name or "Candidate", level=1)
    for paragraph in customization.cover_letter_text.split("\n\n"):
        if paragraph.strip():
            document.add_paragraph(paragraph.strip())
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
