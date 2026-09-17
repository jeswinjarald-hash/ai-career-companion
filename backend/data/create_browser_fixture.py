from pathlib import Path
from docx import Document

output = Path(__file__).parent / "browser_text_resume.docx"
document = Document()
document.add_paragraph("Browser Candidate")
document.add_paragraph("PROFESSIONAL SUMMARY")
document.add_paragraph("Backend developer")
document.add_paragraph("TECHNICAL SKILLS")
document.add_paragraph("Python Fast API")
document.add_paragraph("PROJECTS")
document.add_paragraph("Career Companion API")
document.save(output)
