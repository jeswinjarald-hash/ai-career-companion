"""Resume parser robustness audit: a fixture corpus covering multiple common
student-resume layouts (Part 26 of the audit), PDF/DOCX parity testing (Part 29),
and targeted regression tests for the specific "many false projects / many false
skills" failure modes this audit fixed.

All fixtures are sanitized/synthetic — no real private resume content.
"""

from io import BytesIO

from docx import Document
from pypdf import PdfReader

from app.models import ResumeSection
from app.services.section_detection import detect_resume_sections
from app.services.structured_resume import _projects, build_structured_data
from app.services.text_utils import normalize_resume_text
from pdf_line_fixture import pdf_from_lines
from test_resume_api import create_profile, resume_client


def section(name: str, content: str) -> ResumeSection:
    return ResumeSection(resume_id=1, name=name, original_heading=name.upper(), content=content, position=0)


def structured_from_lines(lines: list[str]) -> dict:
    text = "\n".join(lines)
    detected = detect_resume_sections(text)
    sections = [
        ResumeSection(resume_id=1, name=d.name, original_heading=d.original_heading, content=d.content, position=i)
        for i, d in enumerate(detected)
    ]
    return build_structured_data(sections)


def structured_from_docx(lines: list[str]) -> dict:
    document = Document()
    for line in lines:
        document.add_paragraph(line)
    buffer = BytesIO()
    document.save(buffer)
    from app.services.docx_extraction import _iter_blocks
    from docx.text.paragraph import Paragraph

    buffer.seek(0)
    doc = Document(buffer)
    blocks: list[str] = []
    for block in _iter_blocks(doc):
        if isinstance(block, Paragraph):
            if block.text.strip() or blocks:
                blocks.append(block.text)
    while blocks and not blocks[-1].strip():
        blocks.pop()
    raw_text = "\n".join(blocks)
    normalized = normalize_resume_text(raw_text)
    detected = detect_resume_sections(normalized)
    sections = [
        ResumeSection(resume_id=1, name=d.name, original_heading=d.original_heading, content=d.content, position=i)
        for i, d in enumerate(detected)
    ]
    return build_structured_data(sections)


def structured_from_pdf(lines: list[str]) -> dict:
    pdf_bytes = pdf_from_lines(lines)
    raw_text = PdfReader(BytesIO(pdf_bytes)).pages[0].extract_text()
    normalized = normalize_resume_text(raw_text)
    detected = detect_resume_sections(normalized)
    sections = [
        ResumeSection(resume_id=1, name=d.name, original_heading=d.original_heading, content=d.content, position=i)
        for i, d in enumerate(detected)
    ]
    return build_structured_data(sections)


# ---------------------------------------------------------------------------
# FIXTURE A — "Sam" style: clear headings, bulleted projects, 3 projects.
# ---------------------------------------------------------------------------

FIXTURE_A_LINES = [
    "TECHNICAL SKILLS",
    "Python, Java, JavaScript, SQL, FastAPI, REST API Development, PostgreSQL, MySQL, Scikit-learn, Pandas, Git, GitHub, Postman, VS Code, HTML, CSS",
    "",
    "EDUCATION",
    "B.Tech in Information Technology",
    "Demo Institute of Technology | 2023 - 2027",
    "",
    "EXPERIENCE",
    "Academic Project Developer",
    "- Collaborated with classmates during project development, testing, and version control.",
    "",
    "PROJECTS",
    "AI Career Companion",
    "- Built backend REST APIs using FastAPI and PostgreSQL for candidate profile management.",
    "- Implemented resume upload and structured parsing pipeline.",
    "Technologies: Python, FastAPI, PostgreSQL",
    "",
    "Student Performance Prediction System",
    "- Built a machine learning model to predict student performance using Scikit-learn.",
    "- Cleaned and processed datasets using Pandas.",
    "Technologies: Python, Scikit-learn, Pandas",
    "",
    "Library Management System",
    "- Developed a library management system with issue/return tracking.",
    "- Used MySQL for persistent storage.",
    "Technologies: Java, MySQL",
]


def test_fixture_a_sam_style_produces_three_projects() -> None:
    data = structured_from_lines(FIXTURE_A_LINES)
    assert len(data["projects"]) == 3
    assert [p["title"] for p in data["projects"]] == [
        "AI Career Companion", "Student Performance Prediction System", "Library Management System",
    ]
    assert len(data["education"]) == 1
    assert len(data["experience"]) == 1
    for skill in ("Python", "FastAPI", "PostgreSQL", "Scikit-learn", "MySQL"):
        assert skill in data["skills"]


# ---------------------------------------------------------------------------
# FIXTURE B — "Resume_4" style: no bullets, inline "Name — Built using X, Y, Z",
# long wrapped paragraphs, 2 projects, no experience heading. Also run through a
# realistic (blank-line-losing) PDF extraction, not just raw text.
# ---------------------------------------------------------------------------

FIXTURE_B_LINES = [
    "EDUCATION",
    "Loyola ICAM College of Engineering And Technology, Chennai",
    "B.Tech Information Technology",
    "June 2024 - June 2028",
    "",
    "PROJECTS",
    "Society Finance Management — Built using Node, Express and MongoDB",
    "Developed a full-stack Society Finance Management System independently",
    "using HTML, CSS, JavaScript, Node.js, Express.js, and MongoDB. The project",
    "was designed to manage members, payments, expenses, and loans through a",
    "responsive frontend and REST APIs backed by a MongoDB database with CRUD",
    "operations.",
    "",
    "E-commerce Management — Built using Html, Css, Javascript, React",
    "Developed a dynamic e-commerce web application featuring product listing,",
    "shopping cart, checkout, and an order success page. Implemented",
    "component-based frontend development with interactive state management.",
    "",
    "SKILLS",
    "Frontend Development",
    "Basic Backend Development",
    "SQL Programming",
    "Object-Oriented Programming (OOP)",
    "Problem Solving",
    "",
    "ADDITIONAL",
    "NPTEL Certified in Java Programme",
    "Completed Hands-on Course on Artificial Intelligence and Its Real-Time Application using Python",
    "NPTEL Certified in English",
    "",
    "LANGUAGES",
    "English",
    "Tamil",
]


def test_fixture_b_resume4_style_raw_text() -> None:
    data = structured_from_lines(FIXTURE_B_LINES)
    assert len(data["education"]) == 1
    assert "Loyola ICAM College of Engineering And Technology" in data["education"][0]["raw_text"]
    assert "B.Tech Information Technology" in data["education"][0]["raw_text"]
    assert len(data["projects"]) == 2
    assert [p["title"] for p in data["projects"]] == ["Society Finance Management", "E-commerce Management"]
    assert data["experience"] == []
    for skill in ("Frontend Development", "Basic Backend Development", "SQL", "OOP", "Problem Solving"):
        assert skill in data["skills"]
    assert {entry["raw_text"] for entry in data["languages"]} == {"English", "Tamil"}
    for prose in ("and", "using", "with", "the", "built", "developed", "page", "success", "project", "knowledge", "application"):
        assert prose not in [s.casefold() for s in data["skills"]]


def test_fixture_b_resume4_style_via_realistic_pdf_extraction() -> None:
    # pypdf does not preserve blank-line paragraph gaps as blank text lines (verified
    # empirically), so this exercises the parser with NO blank-line boundary signal
    # at all between the two projects — the harder, more realistic PDF case. The
    # em-dash is swapped for a plain hyphen: this minimal test-PDF builder writes raw
    # UTF-8 bytes, which a simple PDF font's WinAnsi encoding can't represent, a
    # limitation of the test helper, not the app — see the parity test below.
    lines = [line.replace("—", "-") for line in FIXTURE_B_LINES]
    data = structured_from_pdf(lines)
    assert len(data["projects"]) == 2
    assert [p["title"] for p in data["projects"]] == ["Society Finance Management", "E-commerce Management"]
    assert data["experience"] == []
    assert len(data["education"]) == 1


def test_fixture_b_no_prose_words_become_projects() -> None:
    projects = structured_from_lines(FIXTURE_B_LINES)["projects"]
    titles = {p["title"] for p in projects}
    forbidden = {"Society", "Finance", "Management", "Built", "using", "Developed", "a", "full-stack", "and", "The", "project", "was", "designed", "to"}
    assert titles.isdisjoint(forbidden)


# ---------------------------------------------------------------------------
# FIXTURE C — "Project Name\nTechnologies: ..." format.
# ---------------------------------------------------------------------------

def test_fixture_c_project_name_then_technologies_label() -> None:
    lines = [
        "Portfolio Website",
        "Technologies: HTML, CSS, JavaScript",
        "A personal portfolio site showcasing projects and blog posts.",
        "",
        "Weather Dashboard",
        "Technologies: React, Node.js",
        "A weather dashboard consuming a public API with cached results.",
    ]
    projects = _projects([section("projects", "\n".join(lines))])
    assert len(projects) == 2
    assert projects[0]["title"] == "Portfolio Website"
    assert set(projects[0]["technologies"]) >= {"HTML", "CSS", "JavaScript"}
    assert projects[1]["title"] == "Weather Dashboard"
    assert set(projects[1]["technologies"]) >= {"React", "Node.js"}


# ---------------------------------------------------------------------------
# FIXTURE D — "Project Name - React, Node.js, MongoDB" format.
# ---------------------------------------------------------------------------

def test_fixture_d_project_name_dash_technology_list() -> None:
    lines = [
        "Inventory Tracker - React, Node.js, MongoDB",
        "Tracks warehouse stock levels with real-time updates and low-stock alerts.",
        "",
        "Blog Platform - Django, PostgreSQL",
        "A multi-author blogging platform with comments and tagging.",
    ]
    projects = _projects([section("projects", "\n".join(lines))])
    assert len(projects) == 2
    assert projects[0]["title"] == "Inventory Tracker"
    assert set(projects[0]["technologies"]) >= {"React", "Node.js", "MongoDB"}
    assert projects[1]["title"] == "Blog Platform"
    assert set(projects[1]["technologies"]) >= {"Django", "PostgreSQL"}


# ---------------------------------------------------------------------------
# FIXTURE E — resume with a real internship/work-experience section.
# ---------------------------------------------------------------------------

def test_fixture_e_resume_with_real_experience_section() -> None:
    lines = [
        "WORK EXPERIENCE",
        "Backend Developer Intern, BrightPath Labs",
        "Jan 2026 - Jun 2026",
        "- Built and maintained REST APIs serving production traffic.",
        "- Fixed defects reported through the internal bug tracker.",
    ]
    data = structured_from_lines(lines)
    assert len(data["experience"]) == 1
    assert "BrightPath Labs" in data["experience"][0]["raw_text"]


# ---------------------------------------------------------------------------
# FIXTURE F — resume with no projects section at all.
# ---------------------------------------------------------------------------

def test_fixture_f_resume_with_no_projects_section() -> None:
    lines = [
        "EDUCATION",
        "B.Tech Computer Science, Example University, 2027",
        "",
        "SKILLS",
        "Python, SQL, Git",
    ]
    data = structured_from_lines(lines)
    assert data["projects"] == []
    assert len(data["education"]) == 1


# ---------------------------------------------------------------------------
# FIXTURE G — resume with no recognizable skills heading (falls back to
# conservative vocabulary-only scanning of the whole resume).
# ---------------------------------------------------------------------------

def test_fixture_g_resume_with_no_skills_heading_stays_conservative() -> None:
    lines = [
        "EDUCATION",
        "B.Tech Computer Science, Example University, 2027",
        "",
        "PROJECTS",
        "Notes App",
        "- Built with Python and SQLite for local note storage.",
    ]
    data = structured_from_lines(lines)
    assert "Python" in data["skills"]
    assert "SQLite" in data["skills"]
    # No dedicated Skills section -> no unmatched prose word should survive.
    assert "Built" not in data["skills"]
    assert "for" not in [s.casefold() for s in data["skills"]]


# ---------------------------------------------------------------------------
# FIXTURE H — certifications spanning multiple lines stay coherent, grouped
# entries (one per certification), not fragmented into loose words.
# ---------------------------------------------------------------------------

def test_fixture_h_certifications_stay_coherent_multi_word_entries() -> None:
    lines = [
        "CERTIFICATIONS",
        "NPTEL Certified in Java Programming",
        "AWS Cloud Practitioner Essentials - Completed via Coursera",
        "Introduction to Machine Learning with Python",
    ]
    data = structured_from_lines(lines)
    texts = {entry["raw_text"] for entry in data["certifications"]}
    assert "NPTEL Certified in Java Programming" in texts
    assert any("AWS Cloud Practitioner" in text for text in texts)
    assert any("Introduction to Machine Learning" in text for text in texts)
    assert "NPTEL" not in texts
    assert "Java" not in texts


# ---------------------------------------------------------------------------
# PDF vs DOCX parity (Part 29): the same logical resume, both file types,
# should produce materially equivalent structure.
# ---------------------------------------------------------------------------

def test_pdf_and_docx_parity_for_fixture_b() -> None:
    # See the note above: the PDF-path input swaps the em-dash for a plain hyphen to
    # avoid the test PDF builder's WinAnsi-encoding limitation; DOCX handles real
    # Unicode natively, so it uses the original fixture unchanged.
    pdf_lines = [line.replace("—", "-") for line in FIXTURE_B_LINES]
    pdf_data = structured_from_pdf(pdf_lines)
    docx_data = structured_from_docx(FIXTURE_B_LINES)

    assert len(pdf_data["projects"]) == len(docx_data["projects"]) == 2
    assert {p["title"] for p in pdf_data["projects"]} == {p["title"] for p in docx_data["projects"]}
    assert len(pdf_data["education"]) == len(docx_data["education"]) == 1
    assert pdf_data["experience"] == docx_data["experience"] == []
    pdf_skills = set(pdf_data["skills"])
    docx_skills = set(docx_data["skills"])
    assert {"Frontend Development", "Basic Backend Development", "SQL", "OOP", "Problem Solving"} <= pdf_skills
    assert {"Frontend Development", "Basic Backend Development", "SQL", "OOP", "Problem Solving"} <= docx_skills


# ---------------------------------------------------------------------------
# Targeted regressions for the specific "many false projects" / "many false
# skills" failure modes this audit fixed.
# ---------------------------------------------------------------------------

def test_unbulleted_bullets_with_uncommon_action_verbs_do_not_explode_project_count() -> None:
    # None of these verbs are in the continuation-starter stopword list — this must
    # be caught by structural signals (length gate + content-less-fragment merge),
    # not by ever-growing the stopword list.
    lines = [
        "Society Finance Management",
        "Automated payment reconciliation workflows for club members",
        "Streamlined loan approval and disbursement tracking",
        "Integrated MongoDB schema for financial transaction records",
        "Optimized REST endpoints for high-volume requests",
        "Configured role-based access control for admin users",
        "Refactored frontend components for responsiveness",
        "Achieved sub-200ms average API response times",
        "Presented the project at a college tech symposium",
        "",
        "E-commerce Management",
        "Engineered a product catalog with filtering and search",
        "Coordinated checkout and payment gateway integration",
        "Enhanced cart persistence across sessions",
        "Migrated static assets to a CDN for faster loads",
    ]
    projects = _projects([section("projects", "\n".join(lines))])
    assert len(projects) == 2
    assert [p["title"] for p in projects] == ["Society Finance Management", "E-commerce Management"]


def test_leaked_prose_in_skills_section_triggers_sanity_fallback() -> None:
    lines = [
        "SKILLS",
        "Frontend Development, Basic Backend Development, SQL Programming, OOP, Problem Solving",
        "I am a passionate, driven, and enthusiastic student who enjoys learning new things, "
        "building projects, working with teams, solving problems, exploring new technologies, "
        "reading books, playing sports, volunteering, mentoring juniors, and attending workshops, "
        "seminars, and hackathons across the city and beyond every single year",
        "I also like traveling, cooking, photography, painting, writing, chess, cricket, "
        "badminton, swimming, cycling, hiking, camping, gardening, music, dancing, singing, "
        "gaming, and coding in my free time whenever possible",
    ]
    data = structured_from_lines(lines)
    # The contaminated list must not be trusted verbatim — falls back to
    # vocabulary-only matches, staying small and free of prose fragments.
    assert len(data["skills"]) < 10
    for word in ("passionate", "driven", "enthusiastic", "volunteering", "hiking", "cooking"):
        assert word not in [s.casefold() for s in data["skills"]]


# ---------------------------------------------------------------------------
# Parser sanity warnings (Part 21/22): non-destructive diagnostics persisted
# alongside the structured data, never silently blocking a legitimate resume.
# ---------------------------------------------------------------------------

def test_clean_resume_produces_no_warnings() -> None:
    data = structured_from_lines(FIXTURE_A_LINES)
    from app.services.structured_resume import _parser_warnings
    assert _parser_warnings(data) == []


def test_pathological_output_produces_named_warnings() -> None:
    from app.services.structured_resume import _parser_warnings

    pathological = {
        "projects": [{"title": "in", "description": "", "technologies": [], "raw_text": "in"}] * 20,
        "skills": ["and", "using", "SQL"] + [f"Skill{i}" for i in range(30)],
        "experience": [{"raw_text": "in"}],
        "internships": [],
        "education": [{"raw_text": "Loy"}],
    }
    warnings = _parser_warnings(pathological)
    joined = " ".join(warnings)
    assert "suspicious_project_count" in joined
    assert "many_single_word_project_titles" in joined
    assert "suspicious_stopword_project_titles" in joined
    assert "excessive_skill_count" in joined
    assert "excessive_unrecognized_skills" in joined
    assert "suspicious_experience_entry" in joined
    assert "truncated_education_entry" in joined


def test_structure_resume_persists_warnings_in_structured_data(resume_client) -> None:
    from docx import Document

    document = Document()
    for line in FIXTURE_A_LINES:
        document.add_paragraph(line)
    buffer = BytesIO()
    document.save(buffer)

    client, _ = resume_client
    profile_id = create_profile(client)
    upload = client.post(
        f"/api/profiles/{profile_id}/resumes",
        files={"file": ("resume.docx", buffer.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    resume_id = upload.json()["id"]
    client.post(f"/api/resumes/{resume_id}/extract-text")
    client.post(f"/api/resumes/{resume_id}/detect-sections")
    structured = client.post(f"/api/resumes/{resume_id}/structure")

    assert structured.status_code == 200
    assert structured.json()["data"]["parser_warnings"] == []
