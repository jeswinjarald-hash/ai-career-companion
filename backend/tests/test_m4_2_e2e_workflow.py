"""Milestone 4.2 — complete backend end-to-end workflow through the real HTTP API.

Student profile -> resume upload -> parsing -> candidate context -> opportunity
retrieval / matching -> skill gap -> resume customization + cover letter -> interview
preparation -> application tracker -> linked materials -> status/date updates ->
summary + reminders, followed by a cross-user isolation check.

Uses a temporary database and resume directory (the shared `resume_client` fixture),
a generated non-personal DOCX resume, and the suite-wide null LLM guard
(tests/conftest.py). Assertions check cross-module facts — the same user, resume and
job flowing through every step — not just HTTP status codes.
"""

from datetime import datetime, timedelta, timezone
from io import BytesIO

from docx import Document

from app.services.job_dataset_service import load_job_postings
from test_resume_api import register_session, resume_client  # noqa: F401

JOBS = {job.job_id: job for job in load_job_postings()}
DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
ABSENT_SKILLS = {"Kubernetes", "AWS", "Terraform"}


def synthetic_resume() -> bytes:
    document = Document()
    for heading, body in [
        ("PROFESSIONAL SUMMARY", "Backend developer focused on Python APIs"),
        ("TECHNICAL SKILLS", "Python, Fast API, SQL, Postgres, Git, REST APIs"),
        ("EDUCATION", "B.Tech Computer Science, Example Institute of Technology, 2026"),
        ("EXPERIENCE", "Backend Intern at Example Labs, Jan 2025 - Jun 2025. Built REST APIs with FastAPI and PostgreSQL."),
        ("PROJECTS", "Library API - Built a REST API using Python, FastAPI and PostgreSQL"),
    ]:
        document.add_paragraph(heading)
        document.add_paragraph(body)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def ok(response, status: int = 200) -> dict | list:
    assert response.status_code == status, response.text
    return response.json() if response.content else {}


def test_complete_student_workflow_from_profile_to_application_tracking(resume_client) -> None:  # noqa: F811
    client, storage_dir = resume_client
    today = datetime.now(timezone.utc).date()

    # 1. Register + profile
    session = register_session(client, email="e2e-student@example.com", full_name="Eva Endtoend")
    user_id, profile_id = session["user"]["id"], session["profile"]["id"]
    profile = ok(client.patch(f"/api/profiles/{profile_id}", json={"target_roles": ["Backend Developer Intern"], "career_interests": ["backend"]}))
    assert profile["id"] == profile_id and profile["target_roles"] == ["Backend Developer Intern"]

    # 2. Resume upload belongs to this user's profile
    resume = ok(client.post(f"/api/profiles/{profile_id}/resumes", files={"file": ("resume.docx", synthetic_resume(), DOCX_TYPE)}), 201)
    resume_id = resume["id"]
    assert [r["id"] for r in ok(client.get(f"/api/profiles/{profile_id}/resumes"))] == [resume_id]
    assert len(list(storage_dir.iterdir())) == 1

    # 3. Parsing -> structured data -> candidate context
    ok(client.post(f"/api/resumes/{resume_id}/extract-text"))
    ok(client.post(f"/api/resumes/{resume_id}/detect-sections"))
    structured = ok(client.post(f"/api/resumes/{resume_id}/structure"))
    skills = set(structured["data"]["skills"])
    assert {"Python", "FastAPI", "SQL", "PostgreSQL", "Git"} <= skills
    assert not skills & ABSENT_SKILLS
    context = ok(client.post(f"/api/profiles/{profile_id}/candidate-context", params={"resume_id": resume_id}))
    assert context["resume_id"] == resume_id and context["candidate_profile_id"] == profile_id

    # 4. Retrieval + matching return canonical jobs; pick the target from the matches
    search = ok(client.get("/api/jobs/search", params={"q": "FastAPI Python backend internship", "top_k": 5}))
    assert search and all(item["job_id"] in JOBS for item in search)
    matches = ok(client.get(f"/api/resumes/{resume_id}/job-matches", params={"top_k": 10}))
    assert matches and all(m["job_id"] in JOBS for m in matches)
    target = next(m for m in matches if JOBS[m["job_id"]].domain == "Python Backend")
    job_id, job = target["job_id"], JOBS[target["job_id"]]
    assert ok(client.get(f"/api/jobs/{job_id}"))["job_title"] == job.job_title == target["job_title"]
    assert not set(target["matched_required_skills"]) & ABSENT_SKILLS

    # 5. Skill gap for the same resume + job
    gap = ok(client.post(f"/api/resumes/{resume_id}/skill-gap", params={"job_id": job_id}))
    assert (gap["job_id"], gap["job_title"], gap["company"], gap["resume_id"], gap["profile_id"]) == (job_id, job.job_title, job.company, resume_id, profile_id)
    assert not {s["requirement"] for s in gap["strengths"]} & ABSENT_SKILLS

    # 6. Customization + cover letter for the same resume + job
    customization = ok(client.post(f"/api/resumes/{resume_id}/application-customizations", params={"job_id": job_id}))
    assert (customization["job_id"], customization["resume_id"], customization["company"]) == (job_id, resume_id, job.company)
    assert customization["generation"]["mode"] == "deterministic_fallback"
    assert customization["cover_letter_text"].strip() and customization["cover_letter"]
    assert not set(customization["tailored_resume"]["skills"]) & ABSENT_SKILLS
    evidence_ids = {record["evidence_id"] for record in customization["evidence"]}
    assert all(set(sentence["evidence_ids"]) <= evidence_ids for sentence in customization["cover_letter"])
    listed = ok(client.get(f"/api/resumes/{resume_id}/application-customizations", params={"job_id": job_id}))
    assert [item["id"] for item in listed] == [customization["id"]]

    # 7. Interview preparation for the same resume + job
    prep = ok(client.post(f"/api/resumes/{resume_id}/interview-preparations", params={"job_id": job_id}))
    assert (prep["job_id"], prep["resume_id"], prep["job_title"]) == (job_id, resume_id, job.job_title)
    assert prep["questions"] and prep["revision_plan"]

    # 8. Track the opportunity: the backend snapshots the canonical posting
    application = ok(client.post("/api/applications", json={"job_id": job_id, "status": "planning", "deadline": (today + timedelta(days=3)).isoformat()}), 201)
    assert (application["company"], application["job_title"], application["source"]) == (job.company, job.job_title, "dataset")
    assert ok(client.post("/api/applications", json={"job_id": job_id}), 409) is not None

    # 9. Link the generated materials produced above
    linked = ok(client.patch(f"/api/applications/{application['id']}", json={"customization_id": customization["id"], "interview_preparation_id": prep["id"]}))
    assert (linked["customization_id"], linked["interview_preparation_id"]) == (customization["id"], prep["id"])

    # 10. Dates + interview + follow-up, then summary/reminders reflect them
    interview_at = (datetime.now(timezone.utc) + timedelta(days=5)).replace(microsecond=0)
    ok(client.patch(f"/api/applications/{application['id']}", json={
        "follow_up_date": (today + timedelta(days=1)).isoformat(), "interview_at": interview_at.isoformat(), "interview_status": "scheduled", "notes": "Prepare the Library API demo.",
    }))
    summary = ok(client.get("/api/applications/summary"))
    assert (summary["total_applications"], summary["active_applications"], summary["upcoming_deadlines"], summary["interviews_scheduled"]) == (1, 1, 1, 1)
    reminders = ok(client.get("/api/applications/reminders", params={"days": 7}))["reminders"]
    assert {(r["type"], r["application_id"]) for r in reminders} == {("deadline", application["id"]), ("follow_up", application["id"]), ("interview", application["id"])}

    # 11. Status update persists and changes the summary (deadline no longer pending)
    applied = ok(client.patch(f"/api/applications/{application['id']}", json={"status": "applied", "applied_date": today.isoformat()}))
    assert applied["status"] == "applied" and applied["status_updated_at"] > application["status_updated_at"]
    assert ok(client.get(f"/api/applications/{application['id']}"))["status"] == "applied"
    summary = ok(client.get("/api/applications/summary"))
    assert (summary["upcoming_deadlines"], summary["active_applications"], summary["status_counts"]["applied"]) == (0, 1, 1)

    # 12. User isolation: a second student cannot read or reuse anything above
    register_session(client, email="e2e-other@example.com", full_name="Other Student")
    assert client.get(f"/api/resumes/{resume_id}").status_code == 404
    assert client.get(f"/api/resumes/{resume_id}/application-customizations/{customization['id']}").status_code == 404
    assert client.get(f"/api/resumes/{resume_id}/interview-preparations/{prep['id']}").status_code == 404
    assert client.get(f"/api/applications/{application['id']}").status_code == 404
    assert ok(client.get("/api/applications")) == []
    own = ok(client.post("/api/applications", json={"job_id": job_id}), 201)
    assert client.patch(f"/api/applications/{own['id']}", json={"customization_id": customization["id"]}).status_code == 404
    assert client.patch(f"/api/applications/{own['id']}", json={"interview_preparation_id": prep["id"]}).status_code == 404
    assert ok(client.get("/api/applications/summary"))["total_applications"] == 1

    # The first student's data is untouched by the second student's activity.
    client.post("/api/auth/logout")
    ok(client.post("/api/auth/login", json={"email": "e2e-student@example.com", "password": "Password123"}))
    mine = ok(client.get(f"/api/applications/{application['id']}"))
    assert (mine["status"], mine["customization_id"], mine["interview_preparation_id"]) == ("applied", customization["id"], prep["id"])
    assert ok(client.get("/api/auth/me"))["user"]["id"] == user_id
