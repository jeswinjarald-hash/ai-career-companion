from io import BytesIO

from docx import Document
from fastapi.testclient import TestClient

from test_resume_api import create_profile, resume_client


def resume_docx() -> bytes:
    document = Document()
    for heading, body in [
        ("PROFESSIONAL SUMMARY", "Backend developer"),
        ("TECHNICAL SKILLS", "Python, Fast API, Postgres, Docker"),
        ("EDUCATION", "B.Tech Information Technology, Example University, 2024"),
        ("EXPERIENCE", "Backend Intern at Example Co, Jan 2024 - Jun 2024"),
        ("PROJECTS", "Career Companion API"),
    ]:
        document.add_paragraph(heading)
        document.add_paragraph(body)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def prepare(client: TestClient) -> tuple[int, int]:
    profile_id = create_profile(client)
    upload = client.post(
        f"/api/profiles/{profile_id}/resumes",
        files={
            "file": (
                "resume.docx", resume_docx(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert upload.status_code == 201
    resume_id = upload.json()["id"]
    assert client.post(f"/api/resumes/{resume_id}/extract-text").status_code == 200
    assert client.post(f"/api/resumes/{resume_id}/detect-sections").status_code == 200
    return profile_id, resume_id


def test_structured_resume_and_candidate_context_are_persisted_and_idempotent(resume_client) -> None:
    client, _ = resume_client
    profile_id, resume_id = prepare(client)

    first = client.post(f"/api/resumes/{resume_id}/structure")
    second = client.post(f"/api/resumes/{resume_id}/structure")
    assert first.status_code == 200
    assert second.status_code == 200
    assert first.json()["id"] == second.json()["id"]
    data = second.json()["data"]
    assert "Python" in data["skills"]
    assert "FastAPI" in data["skills"]
    assert "PostgreSQL" in data["skills"]
    assert data["education"][0]["degree"] == "B.Tech"

    context = client.post(f"/api/profiles/{profile_id}/candidate-context?resume_id={resume_id}")
    context_again = client.post(f"/api/profiles/{profile_id}/candidate-context?resume_id={resume_id}")
    assert context.status_code == 200
    assert context_again.status_code == 200
    assert context.json()["id"] == context_again.json()["id"]
    assert "FastAPI" in context.json()["context_json"]["combined"]["skills"]
    assert client.get(f"/api/profiles/{profile_id}/candidate-context?resume_id={resume_id}").status_code == 200


def test_structure_and_context_require_previous_stages(resume_client) -> None:
    client, _ = resume_client
    profile_id = create_profile(client)
    response = client.post(f"/api/profiles/{profile_id}/candidate-context?resume_id=999")
    assert response.status_code == 404
    assert client.post("/api/resumes/999/structure").status_code == 404


def resume_docx_with_grouped_content() -> bytes:
    document = Document()
    for heading, body_lines in [
        ("PROFESSIONAL SUMMARY", ["Backend developer"]),
        ("TECHNICAL SKILLS", ["Python, FastAPI, Postgres, Docker"]),
        ("EDUCATION", [
            "B.Tech in Information Technology - Example University",
            "Expected Graduation: 2027 | CGPA: 8.4/10",
        ]),
        ("PROJECTS", [
            "Career Companion API",
            " Built REST APIs for candidate profile creation and resume upload.",
            " Implemented PostgreSQL-backed persistence and validation.",
            "Second Project - Data Pipeline",
            " Cleaned and transformed data using Pandas and NumPy.",
        ]),
    ]:
        document.add_paragraph(heading)
        for line in body_lines:
            document.add_paragraph(line)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def test_reprocessing_upserts_corrected_grouping_without_duplicating(resume_client) -> None:
    client, _ = resume_client
    profile_id = create_profile(client)
    upload = client.post(
        f"/api/profiles/{profile_id}/resumes",
        files={
            "file": (
                "resume.docx", resume_docx_with_grouped_content(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    resume_id = upload.json()["id"]
    assert client.post(f"/api/resumes/{resume_id}/extract-text").status_code == 200
    assert client.post(f"/api/resumes/{resume_id}/detect-sections").status_code == 200

    first = client.post(f"/api/resumes/{resume_id}/structure")
    assert first.status_code == 200
    first_data = first.json()

    # These counts are naturally derived from the input content by the real
    # extraction/section-detection/structuring pipeline, not hardcoded.
    assert len(first_data["data"]["projects"]) == 2
    assert len(first_data["data"]["education"]) == 1
    assert first_data["data"]["education"][0]["graduation_year"] == 2027

    # Re-running structuring (what the "Reprocess this resume" action, or a
    # parser fix landing after the first run, would trigger) must upsert the
    # existing StructuredResume row rather than create a second one.
    second = client.post(f"/api/resumes/{resume_id}/structure")
    assert second.status_code == 200
    assert second.json()["id"] == first_data["id"]
    assert len(second.json()["data"]["projects"]) == 2

    stored = client.get(f"/api/resumes/{resume_id}/structured")
    assert stored.status_code == 200
    assert stored.json()["id"] == first_data["id"]

    # Reprocessing the candidate context (what the reprocess action also does)
    # likewise upserts rather than duplicating.
    context_first = client.post(f"/api/profiles/{profile_id}/candidate-context?resume_id={resume_id}")
    context_second = client.post(f"/api/profiles/{profile_id}/candidate-context?resume_id={resume_id}")
    assert context_first.status_code == 200
    assert context_second.status_code == 200
    assert context_first.json()["id"] == context_second.json()["id"]
