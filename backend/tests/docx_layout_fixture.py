"""Synthetic DOCX resumes in the layouts Word templates commonly use: contact
details in the page header, and a sidebar column held in a floating text box
(stored as ``w:txbxContent`` inside a drawing, with a VML fallback copy of the same
text for older readers)."""

from io import BytesIO

from docx import Document
from docx.oxml import parse_xml

_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _textbox_xml(lines: list[str]) -> str:
    paragraphs = "".join(f'<w:p><w:r><w:t xml:space="preserve">{line}</w:t></w:r></w:p>' for line in lines)
    return (
        f'<w:r xmlns:w="{_W}" xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
        'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
        'xmlns:v="urn:schemas-microsoft-com:vml">'
        "<mc:AlternateContent><mc:Choice Requires=\"wps\"><w:drawing><wp:anchor><a:graphic><a:graphicData>"
        f"<wps:wsp><wps:txbx><w:txbxContent>{paragraphs}</w:txbxContent></wps:txbx></wps:wsp>"
        "</a:graphicData></a:graphic></wp:anchor></w:drawing></mc:Choice>"
        f"<mc:Fallback><w:pict><v:shape><v:textbox><w:txbxContent>{paragraphs}</w:txbxContent></v:textbox></v:shape></w:pict></mc:Fallback>"
        "</mc:AlternateContent></w:r>"
    )


def docx_with_header_and_sidebar(header_lines: list[str], body_lines: list[str], sidebar_lines: list[str]) -> bytes:
    document = Document()
    header = document.sections[0].header
    header.paragraphs[0].text = header_lines[0]
    for line in header_lines[1:]:
        header.add_paragraph(line)
    anchor = document.add_paragraph()
    anchor._p.append(parse_xml(_textbox_xml(sidebar_lines)))
    for line in body_lines:
        document.add_paragraph(line)
    output = BytesIO()
    document.save(output)
    return output.getvalue()
