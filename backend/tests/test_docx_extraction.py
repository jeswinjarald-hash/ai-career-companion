from io import BytesIO

from docx import Document
from fastapi.testclient import TestClient

from test_resume_api import create_profile, resume_client


def docx_bytes(with_table: bool = True, with_text: bool = True) -> bytes:
    document = Document()
    if with_text:
        document.add_paragraph("Candidate Name")
        document.add_paragraph("Summary")
        document.add_paragraph("Python Developer")
    if with_table:
        table = document.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "Skill"
        table.rows[0].cells[1].text = "Level"
        row = table.add_row().cells
        row[0].text = "Python"
        row[1].text = "Advanced"
        row = table.add_row().cells
        row[0].text = "SQL"
        row[1].text = "Intermediate"
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def upload_docx(client: TestClient, profile_id: int, contents: bytes):
    response = client.post(
        f"/api/profiles/{profile_id}/resumes",
        files={
            "file": (
                "resume.docx",
                contents,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert response.status_code == 201, response.text
    return response.json()


def test_extracts_docx_paragraphs_and_table_content(resume_client) -> None:
    client, _ = resume_client
    resume = upload_docx(client, create_profile(client), docx_bytes())

    response = client.post(f"/api/resumes/{resume['id']}/extract-text")
    assert response.status_code == 200, response.text
    extraction = response.json()
    assert extraction["status"] == "text_extracted"
    assert extraction["page_count"] == 0
    for expected in ["Candidate Name", "Summary", "Python Developer", "Skill | Level", "Python | Advanced", "SQL | Intermediate"]:
        assert expected in extraction["normalized_text"]
    assert client.get(f"/api/resumes/{resume['id']}").json()["status"] == "text_extracted"


def test_docx_extraction_is_idempotent(resume_client) -> None:
    client, _ = resume_client
    resume = upload_docx(client, create_profile(client), docx_bytes(with_table=False))

    first = client.post(f"/api/resumes/{resume['id']}/extract-text")
    second = client.post(f"/api/resumes/{resume['id']}/extract-text")
    assert first.status_code == 200
    assert second.status_code == 200
    assert first.json()["id"] == second.json()["id"]
    assert first.json()["normalized_text"] == second.json()["normalized_text"]


def test_blank_paragraphs_between_content_survive_as_blank_lines(resume_client) -> None:
    # A blank paragraph between two paragraphs of real content is the author's own
    # visual separator (e.g. between two resume projects) — it must survive
    # extraction as a blank line, not be silently dropped. Leading/trailing blank
    # paragraphs (common Word formatting noise) are still trimmed away.
    document = Document()
    document.add_paragraph("")
    document.add_paragraph("First Entry")
    document.add_paragraph("Detail one.")
    document.add_paragraph("")
    document.add_paragraph("Second Entry")
    document.add_paragraph("Detail two.")
    document.add_paragraph("")
    output = BytesIO()
    document.save(output)

    client, _ = resume_client
    resume = upload_docx(client, create_profile(client), output.getvalue())

    response = client.post(f"/api/resumes/{resume['id']}/extract-text")
    assert response.status_code == 200, response.text
    normalized = response.json()["normalized_text"]

    assert "First Entry\nDetail one.\n\nSecond Entry\nDetail two." in normalized
    assert not normalized.startswith("\n")
    assert not normalized.endswith("\n")


def test_empty_docx_fails_cleanly(resume_client) -> None:
    client, _ = resume_client
    resume = upload_docx(client, create_profile(client), docx_bytes(with_table=False, with_text=False))

    response = client.post(f"/api/resumes/{resume['id']}/extract-text")
    assert response.status_code == 400
    assert response.json()["detail"] == "No extractable text was found in this DOCX resume."
    extraction = client.get(f"/api/resumes/{resume['id']}/extraction").json()
    assert extraction["status"] == "extraction_failed"
    assert extraction["error_message"] == response.json()["detail"]