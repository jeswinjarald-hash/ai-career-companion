"""Writes a synthetic DOCX resume for the Playwright upload journey.

Every detail is fictional (the person, the e-mail on the reserved example.com domain,
the phone number in the 555 fictional range, the college and the projects). It
contains no real personal information and is regenerated on every run, so nothing
binary is committed.
"""

import sys
from pathlib import Path

from docx import Document

# Backend / data profile (used by most journeys).
LINES: list[tuple[str, str]] = [
    ("title", "Alex Sample"),
    ("text", "alex.sample@example.com | +1 555 0100 | Sampletown"),
    ("heading", "Summary"),
    ("text", "Computer science student interested in backend development and data analysis."),
    ("heading", "Skills"),
    ("text", "Python, SQL, FastAPI, React, TypeScript, Git, Pandas, Machine Learning"),
    ("heading", "Education"),
    ("text", "B.Tech in Computer Science, Example Institute of Technology, 2022 - 2026"),
    ("heading", "Projects"),
    ("text", "Library Catalogue API"),
    ("bullet", "Built a REST API with Python and FastAPI backed by a SQL database."),
    ("text", "Sales Data Dashboard"),
    ("bullet", "Analysed sample sales data with Pandas and visualised trends in a React dashboard."),
    ("heading", "Experience"),
    ("text", "Software Engineering Intern, Example Labs, June 2025 - August 2025"),
    ("bullet", "Wrote Python scripts to automate data validation and reporting."),
    ("heading", "Certifications"),
    ("text", "Introduction to Machine Learning (online course)"),
]


# A second, deliberately different domain, so tests that switch between two resumes
# get a different top job match for each.
MARKETING_LINES: list[tuple[str, str]] = [
    ("title", "Jordan Sample"),
    ("text", "jordan.sample@example.com | +1 555 0101 | Sampletown"),
    ("heading", "Summary"),
    ("text", "Business student focused on digital marketing, social media and content strategy."),
    ("heading", "Skills"),
    ("text", "Digital Marketing, SEO, Social Media Marketing, Content Writing, Google Analytics, Canva, Email Marketing"),
    ("heading", "Education"),
    ("text", "BBA in Marketing, Example College of Business, 2022 - 2025"),
    ("heading", "Projects"),
    ("text", "Campus Event Campaign"),
    ("bullet", "Planned a social media campaign and tracked engagement with Google Analytics."),
    ("heading", "Experience"),
    ("text", "Marketing Intern, Example Retail, May 2024 - July 2024"),
    ("bullet", "Wrote blog posts and email newsletters and improved on-page SEO."),
]


def build(path: Path, lines: list[tuple[str, str]] = LINES) -> None:
    document = Document()
    for kind, text in lines:
        if kind == "title":
            document.add_heading(text, level=0)
        elif kind == "heading":
            document.add_heading(text, level=1)
        elif kind == "bullet":
            document.add_paragraph(text, style="List Bullet")
        else:
            document.add_paragraph(text)
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))


if __name__ == "__main__":
    build(Path(sys.argv[1]), MARKETING_LINES if len(sys.argv) > 2 and sys.argv[2] == "marketing" else LINES)
